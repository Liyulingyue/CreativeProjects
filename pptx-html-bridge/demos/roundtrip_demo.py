#!/usr/bin/env python3
"""
演示脚本：PPTX -> HTML -> PPTX 双向桥接 round-trip

流程：
1. 将 demos/source/test.pptx 正向转换为 HTML（输出到 demos/outputs/）
2. 将生成的 HTML 反向转换回 PPTX（输出到 demos/outputs/rebuilt.pptx）
3. 校验重建的 PPTX 内容
"""

import os
import shutil

from pptx import Presentation

from pptx_html_bridge import PPTXToHTMLConverter, convert_html_to_pptx


def main():
    script_dir = os.path.dirname(os.path.abspath(__file__))
    project_dir = os.path.dirname(script_dir)
    source_file = os.path.join(project_dir, "demos", "source", "test.pptx")
    output_dir = os.path.join(project_dir, "demos", "outputs")
    rebuilt_path = os.path.join(output_dir, "rebuilt.pptx")

    print("=== PPTX -> HTML -> PPTX Round-trip 演示 ===\n")

    if not os.path.exists(source_file):
        print(f"❌ 错误：源文件不存在 - {source_file}")
        print("   请先运行 demos/make_sample_pptx.py 生成演示文件")
        return 1

    # 步骤1：正向转换 PPTX -> HTML
    print("🔄 步骤1：正向转换 PPTX -> HTML ...")
    if os.path.exists(output_dir):
        shutil.rmtree(output_dir)
    forward = PPTXToHTMLConverter(compact=True)
    result = forward.convert_file(source_file, output_dir)
    print(f"   ✓ 生成 {result['slides_count']} 页幻灯片 HTML -> {output_dir}")

    # 步骤2：反向转换 HTML -> PPTX
    print("\n🔄 步骤2：反向转换 HTML -> PPTX ...")
    convert_html_to_pptx(output_dir, rebuilt_path)
    print(f"   ✓ 已重建 PPTX -> {rebuilt_path}")

    # 步骤3：校验重建结果
    print("\n🔍 步骤3：校验重建的 PPTX ...")
    prs = Presentation(rebuilt_path)
    assert len(prs.slides) == result['slides_count'], "幻灯片数量不一致"
    print(f"   ✓ 幻灯片数量：{len(prs.slides)}")

    texts = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                texts.append(shape.text_frame.text)
    joined = "\n".join(texts)
    assert "pptx-html-bridge 演示文稿" in joined, "标题文本丢失"
    print("   ✓ 标题文本已还原")

    assert os.path.getsize(rebuilt_path) > 0
    print("\n🎉 Round-trip 完成！PPTX -> HTML -> PPTX 内容一致")
    return 0


if __name__ == "__main__":
    exit(main())

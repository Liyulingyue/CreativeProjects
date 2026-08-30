"""Tests for advanced features: groups, theme colors, bullets, hyperlinks,
notes, rotation, table styling, and directory conversion."""

import os
import shutil

import pytest

from pptx import Presentation

from pptx_html_bridge import PPTXToHTMLConverter, convert_pptx_directory, convert_html_to_pptx
from pptx_html_bridge.html_parsers import parse_slide_html
from pptx_html_bridge.themes import resolve_theme_color


@pytest.fixture
def advanced_html(sample_pptx, tmp_path):
    output_dir = os.path.join(str(tmp_path), "html")
    PPTXToHTMLConverter().convert_file(sample_pptx, output_dir)
    return output_dir


@pytest.fixture
def advanced_pptx(sample_pptx):
    return Presentation(sample_pptx)


# ---------- forward conversion ----------

def test_group_shapes_rendered(advanced_html):
    html = open(os.path.join(advanced_html, "slides", "slide4.html"), encoding="utf-8").read()
    assert "in group A" in html
    assert "in group B" in html
    # both group children got computed absolute positions
    assert html.count("in group") == 2


def test_theme_color_text(advanced_html, advanced_pptx):
    html = open(os.path.join(advanced_html, "slides", "slide4.html"), encoding="utf-8").read()
    slide4 = advanced_pptx.slides[3]
    expected = None
    for shape in slide4.shapes:
        if shape.has_text_frame and "themed text" in shape.text_frame.text:
            run = shape.text_frame.paragraphs[0].runs[0]
            expected = resolve_theme_color(advanced_pptx, run.font.color.theme_color)
    assert expected is not None
    assert f"color: {expected}" in html


def test_bullets_rendered(advanced_html):
    html = open(os.path.join(advanced_html, "slides", "slide4.html"), encoding="utf-8").read()
    assert "&#8226;" in html or "\u2022" in html or "•" in html
    assert "margin-left" in html
    assert "1.&nbsp;" in html  # numbered item


def test_hyperlink_rendered(advanced_html):
    html = open(os.path.join(advanced_html, "slides", "slide4.html"), encoding="utf-8").read()
    assert 'href="https://example.com"' in html
    assert "example link" in html


def test_notes_rendered(advanced_html):
    html = open(os.path.join(advanced_html, "slides", "slide4.html"), encoding="utf-8").read()
    assert 'class="notes"' in html
    assert "组合形状与主题色测试" in html


def test_text_rotation_rendered(advanced_html):
    html = open(os.path.join(advanced_html, "slides", "slide4.html"), encoding="utf-8").read()
    assert "rotate(15" in html and "deg)" in html
    assert "rotated label" in html


def test_table_styling_rendered(advanced_html):
    html = open(os.path.join(advanced_html, "slides", "slide4.html"), encoding="utf-8").read()
    assert 'colspan="2"' in html          # merged header
    assert "background-color: #cc2222" in html  # filled cell
    assert "merged header" in html


# ---------- reverse conversion ----------

def test_reverse_notes_and_rotation(advanced_html, tmp_path):
    out = os.path.join(str(tmp_path), "rebuilt.pptx")
    convert_html_to_pptx(advanced_html, out)
    prs = Presentation(out)
    slide4 = prs.slides[3]
    assert slide4.notes_slide.notes_text_frame.text == "备注：组合形状与主题色测试"

    rotations = [sh.rotation for sh in slide4.shapes
                 if sh.has_text_frame and "rotated label" in sh.text_frame.text]
    assert rotations and abs(rotations[0] - 15.0) < 0.1


def test_reverse_hyperlink(advanced_html, tmp_path):
    out = os.path.join(str(tmp_path), "rebuilt.pptx")
    convert_html_to_pptx(advanced_html, out)
    prs = Presentation(out)
    slide4 = prs.slides[3]
    hrefs = []
    for shape in slide4.shapes:
        if shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                for run in p.runs:
                    if run.hyperlink.address:
                        hrefs.append(run.hyperlink.address)
    assert "https://example.com" in hrefs


def test_reverse_table_fill_and_merge(advanced_html, tmp_path):
    out = os.path.join(str(tmp_path), "rebuilt.pptx")
    convert_html_to_pptx(advanced_html, out)
    prs = Presentation(out)
    slide4 = prs.slides[3]
    tables = [sh for sh in slide4.shapes if sh.has_table]
    assert len(tables) == 1
    table = tables[0].table
    assert table.cell(0, 0).text == "merged header"
    assert str(table.cell(1, 1).fill.fore_color.rgb) == "CC2222"


# ---------- directory conversion ----------

def test_directory_conversion_isolated_outputs(sample_pptx, tmp_path):
    src_dir = os.path.join(str(tmp_path), "src")
    os.makedirs(src_dir)
    shutil.copy(sample_pptx, os.path.join(src_dir, "alpha.pptx"))
    shutil.copy(sample_pptx, os.path.join(src_dir, "beta.pptx"))
    out_dir = os.path.join(str(tmp_path), "out")

    result = convert_pptx_directory(src_dir, out_dir)
    assert result["converted_files"] == 2

    # per-file subdirectories, no collisions
    for base in ("alpha", "beta"):
        assert os.path.isfile(os.path.join(out_dir, base, "slides", "slide1.html"))
        assert os.path.isfile(os.path.join(out_dir, base, "media", f"{base}_slide2_img0.png"))
        assert os.path.isfile(os.path.join(out_dir, base, f"{base}_index.html"))

    main_html = open(os.path.join(out_dir, "main.html"), encoding="utf-8").read()
    assert 'href="alpha/alpha_index.html"' in main_html
    assert 'href="beta/beta_index.html"' in main_html


def test_reverse_works_on_directory_output(sample_pptx, tmp_path):
    src_dir = os.path.join(str(tmp_path), "src")
    os.makedirs(src_dir)
    shutil.copy(sample_pptx, os.path.join(src_dir, "alpha.pptx"))
    out_dir = os.path.join(str(tmp_path), "out")
    convert_pptx_directory(src_dir, out_dir)

    rebuilt = os.path.join(str(tmp_path), "alpha_rebuilt.pptx")
    result = convert_html_to_pptx(os.path.join(out_dir, "alpha"), rebuilt)
    assert result["slides_count"] == 4
    prs = Presentation(rebuilt)
    assert len(prs.slides) == 4

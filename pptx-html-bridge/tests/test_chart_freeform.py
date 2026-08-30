"""Tests for chart / freeform / transition rendering and rebuild."""

import os

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from pptx_html_bridge import convert_html_to_pptx


def test_chart_rendered_as_data_table(advanced_html):
    html = open(os.path.join(advanced_html, "slides", "slide5.html"), encoding="utf-8").read()
    assert 'data-chart-type="column"' in html
    assert 'class="chart-data"' in html
    assert "<th>Q1</th>" in html
    assert "<th>sales</th>" in html
    assert "<td>25.0</td>" in html


def test_freeform_rendered_as_svg(advanced_html):
    html = open(os.path.join(advanced_html, "slides", "slide5.html"), encoding="utf-8").read()
    assert 'class="freeform"' in html
    assert "<polygon" in html
    assert "fill: #22aa66" in html


def test_fade_transition_css(advanced_html):
    html = open(os.path.join(advanced_html, "slides", "slide5.html"), encoding="utf-8").read()
    assert "slideFadeIn" in html


def test_chart_and_freeform_rebuilt(advanced_html, tmp_path):
    out = os.path.join(str(tmp_path), "rebuilt.pptx")
    convert_html_to_pptx(advanced_html, out)
    prs = Presentation(out)
    slide5 = prs.slides[4]

    # native chart rebuilt from the data table
    charts = [sh for sh in slide5.shapes if sh.has_chart]
    assert len(charts) == 1
    plot = charts[0].chart.plots[0]
    assert [str(c) for c in plot.categories] == ["Q1", "Q2", "Q3"]
    series = list(plot.series)
    assert len(series) == 1
    assert list(series[0].values) == [12.0, 25.0, 31.0]

    # freeform polygon rebuilt
    freeforms = [sh for sh in slide5.shapes if sh.shape_type == MSO_SHAPE_TYPE.FREEFORM]
    assert len(freeforms) == 1

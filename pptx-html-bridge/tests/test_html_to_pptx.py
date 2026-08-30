"""Tests for the reverse direction: HTML -> PPTX (round trip)."""

import os

import pytest

from pptx import Presentation
from pptx.util import Emu

from pptx_html_bridge import PPTXToHTMLConverter, HTMLToPPTXConverter, convert_html_to_pptx
from pptx_html_bridge.html_parsers import parse_slide_html


@pytest.fixture
def converted_html(sample_pptx, tmp_path):
    """Forward-convert the sample pptx and return the html output dir."""
    output_dir = os.path.join(str(tmp_path), "html")
    PPTXToHTMLConverter(compact=True).convert_file(sample_pptx, output_dir)
    return output_dir


def test_parse_slide_model(converted_html):
    model = parse_slide_html(os.path.join(converted_html, "slides", "slide1.html"))
    assert model["width_px"] == 960
    assert model["height_px"] == 720
    assert model["background_color"] == "1b2a4a"
    kinds = [s["kind"] for s in model["shapes"]]
    assert kinds.count("text") == 2
    # title run: escaped text is unescaped by the HTML parser
    texts = [r["text"] for s in model["shapes"] if s["kind"] == "text"
             for p in s["paragraphs"] for r in p["runs"]]
    assert "Hello & <World>" in texts


def test_parse_slide2_shapes(converted_html):
    model = parse_slide_html(os.path.join(converted_html, "slides", "slide2.html"))
    kinds = sorted(s["kind"] for s in model["shapes"])
    assert "image" in kinds
    assert "autoshape" in kinds
    assert "line" in kinds
    autoshape = next(s for s in model["shapes"] if s["kind"] == "autoshape")
    assert autoshape["fill_color"] == "2563eb"
    line = next(s for s in model["shapes"] if s["kind"] == "line")
    assert line["stroke_color"] == "666666"


def test_parse_slide3_table(converted_html):
    model = parse_slide_html(os.path.join(converted_html, "slides", "slide3.html"))
    table = next(s for s in model["shapes"] if s["kind"] == "table")
    assert table["rows"][0] == ["A&B", "<tag>"]
    assert table["rows"][1] == ["cell", "value"]


def test_html_to_pptx_roundtrip(converted_html, tmp_path):
    output_pptx = os.path.join(str(tmp_path), "rebuilt.pptx")
    result = convert_html_to_pptx(converted_html, output_pptx)

    assert result["slides_count"] == 3
    assert os.path.isfile(output_pptx)

    prs = Presentation(output_pptx)
    assert len(prs.slides) == 3

    # slide size preserved (px -> EMU)
    assert prs.slide_width == Emu(960 * 9525)
    assert prs.slide_height == Emu(720 * 9525)

    # slide 1: white bold title text survived the round trip
    slide1 = prs.slides[0]
    texts = []
    for shape in slide1.shapes:
        if shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                for run in p.runs:
                    texts.append((run.text, run.font.size, run.font.bold))
    assert any(t[0] == "Hello & <World>" for t in texts)
    title = next(t for t in texts if t[0] == "Hello & <World>")
    assert title[2] is True
    assert title[1] is not None and abs(title[1].pt - 40.0) < 0.6

    # slide 2: picture part present
    slide2 = prs.slides[1]
    pic_parts = [s for s in slide2.shapes if s.shape_type == 13]  # PICTURE
    assert len(pic_parts) == 1

    # slide 3: table with escaped cell text unescaped
    slide3 = prs.slides[2]
    tables = [s for s in slide3.shapes if s.has_table]
    assert len(tables) == 1
    assert tables[0].table.cell(0, 0).text == "A&B"


def test_slide1_background_color_roundtrip(converted_html, tmp_path):
    output_pptx = os.path.join(str(tmp_path), "rebuilt.pptx")
    convert_html_to_pptx(converted_html, output_pptx)
    prs = Presentation(output_pptx)
    slide1 = prs.slides[0]
    fill = slide1.background.fill
    assert fill.type is not None
    assert str(fill.fore_color.rgb) == "1B2A4A"


def test_html_to_pptx_missing_dir(tmp_path):
    try:
        convert_html_to_pptx(os.path.join(str(tmp_path), "nope"))
        assert False, "expected FileNotFoundError"
    except FileNotFoundError:
        pass

"""Shared pytest fixtures: builds a small sample PPTX for conversion tests."""

import base64
import os

import pytest

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

PNG_BYTES = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAIAAAACCAYAAABytg0kAAAAEklEQVR4nGP8z8Dwn4GBgYEA"
    "KicDdwTgXbAAAAAAAElFTkSuQmCC"
)


@pytest.fixture
def sample_pptx(tmp_path):
    path = os.path.join(str(tmp_path), "sample.pptx")
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    # Slide 1: dark background + white title + subtitle (no explicit color)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = RGBColor(0x1B, 0x2A, 0x4A)

    title = slide.shapes.add_textbox(Inches(1), Inches(2.5), Inches(8), Inches(1.2))
    p = title.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = "Hello & <World>"
    run.font.size = Pt(40)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    sub = slide.shapes.add_textbox(Inches(1), Inches(3.8), Inches(8), Inches(0.8))
    p = sub.text_frame.paragraphs[0]
    run = p.add_run()
    run.text = "plain subtitle"
    run.font.size = Pt(20)

    # Slide 2: light background + autoshape + line + picture
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.8), Inches(3.5), Inches(2))
    box.fill.solid()
    box.fill.fore_color.rgb = RGBColor(0x25, 0x63, 0xEB)

    line = slide.shapes.add_connector(1, Inches(0.8), Inches(3.4), Inches(9.2), Inches(3.4))
    line.line.color.rgb = RGBColor(0x66, 0x66, 0x66)

    import io
    slide.shapes.add_picture(io.BytesIO(PNG_BYTES), Inches(1), Inches(4), width=Inches(2), height=Inches(2))

    # Slide 3: table
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    table_shape = slide.shapes.add_table(2, 2, Inches(1.5), Inches(1.5), Inches(6), Inches(2))
    table_shape.table.cell(0, 0).text = "A&B"
    table_shape.table.cell(0, 1).text = "<tag>"
    table_shape.table.cell(1, 0).text = "cell"
    table_shape.table.cell(1, 1).text = "value"

    prs.save(path)
    return path
